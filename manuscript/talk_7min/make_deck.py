"""7분 사내 발표회 PPT (Widescreen 13.333 x 7.5 in). 노트에는 SCRIPT_7min_KR_trim.md(7분 맞춤판)이 들어간다.

스타일: 흰 배경, 강조색은 navy 하나, 나머지는 회색. 색 띠·색 상자 대신 얇은 선과 굵은 글씨로 위계를 만든다.
layout helper는 presentation/make_talk_5min.py를 쓴다(slide-design-patterns Preset A의 좌표계).

usage:
    .venv/bin/python manuscript/talk_7min/lab/scenario_lab.py --defaults   # lab_*.png + lab_summary.json
    .venv/bin/python manuscript/talk_7min/make_figs.py
    .venv/bin/python manuscript/talk_7min/make_deck.py
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "presentation"))
import make_talk_5min as d  # noqa: E402
from make_talk_5min import MSO_SHAPE, fit, frame, hline, kpi, rect, text, vline  # noqa: E402
from pptx import Presentation  # noqa: E402

TF = HERE / "figs"
OUT = HERE / "SRAM_Vmin_7min_KR.pptx"
SCRIPT = HERE / "SCRIPT_7min_KR_trim.md"
# 시나리오 수치는 손으로 옮기지 않는다: lab/scenario_lab.py --defaults가 만든 값을 읽는다
SM = json.loads((TF / "lab_summary.json").read_text(encoding="utf-8"))
T = SM["target"]
NEED = {k: v["change"] for k, v in SM["levers"].items()}

ACC = "1428A0"      # 유일한 강조색 (navy)
INK = "1A1A1A"      # 본문
MUTE = "7F7F7F"     # 보조 글씨
LINE = "BFBFBF"     # 상자 테두리
SOFT = "F3F3F3"     # 아주 옅은 바탕 (꼭 필요한 곳만)

TITLE_KR = ["공정 변동 하의 SRAM Vmin 순·역방향 추정을 위한", "physics-guided framework"]
BRAND = "SRAM Vmin 순·역방향 추정"
TITLE_EN = "A Physics-Guided Framework for Forward and Inverse SRAM Vmin Estimation Under Process Variation"


def load_notes() -> dict[int, str]:
    """대본의 '## N. (시간) 제목' 아래 문단을 슬라이드 N의 노트로 쓴다."""
    parts = re.split(r"^## (\d+)\..*$", SCRIPT.read_text(encoding="utf-8"), flags=re.M)
    return {int(n): body.split("\n---")[0].strip() for n, body in zip(parts[1::2], parts[2::2])}


def box(s, x, y, w, h, lines, size=15, color=INK, bold=False, strong=False, align="c", spacing=1.1):
    """흰 상자 + 얇은 테두리. strong=True면 navy 테두리·navy 글씨."""
    rect(s, x, y, w, h, "FFFFFF", line=ACC if strong else LINE, lw=1.75 if strong else 0.75)
    text(s, x + 0.8, y, w - 1.6, h, lines, size=size, color=ACC if strong else color, bold=bold or strong,
         align=align, anchor="m", spacing=spacing)


def message(s, y, msg):
    """장표의 결론 한 줄: 위에 얇은 선, navy 굵은 글씨."""
    hline(s, 3.8, y, 92.4, color=LINE, pt=0.75)
    text(s, 3.8, y + 1.2, 92.4, 7, msg, size=19, color=ACC, bold=True, align="c", anchor="m")


def label(s, x, y, w, t, color=MUTE, size=14, align="l"):
    text(s, x, y, w, 4.5, t, size=size, color=color, bold=True, align=align)


def flow(s, x0, y, w_total, items, hi=None, h=11, size=14):
    """→로 이은 단계 상자. hi 번째 상자만 navy로 강조."""
    n, gap = len(items), 3.2
    w = (w_total - gap * (n - 1)) / n
    for i, body in enumerate(items):
        x = x0 + i * (w + gap)
        box(s, x, y, w, h, body.split("\n"), size=size, strong=(i == hi), spacing=1.0)
        if i < n - 1:
            text(s, x + w, y, gap, h, "→", size=18, color=MUTE, align="c", anchor="m")


# ------------------------------------------------------------------ slides
def cover(prs, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 6, 28, 88, 18, TITLE_KR, size=34, color=INK, bold=True, align="c", anchor="m", spacing=1.15)
    hline(s, 44, 49.5, 12, color=ACC, pt=1.5)
    text(s, 8, 53, 84, 8, TITLE_EN, size=15, color=MUTE, align="c")
    text(s, 5, 78, 90, 4, "[발표자]  ·  [소속]", size=17, color=INK, align="c")
    text(s, 5, 84, 90, 3.5, "2026", size=13, color=MUTE, align="c")
    s.notes_slide.notes_text_frame.text = notes[1]


def background(prs, notes):
    s = frame(prs, 2, "배경: FinFET Low power VOC와 SRAM cell 검토 수요", notes[2],
              foot="세 가지 VOC 중 이 연구는 Vmin을 다룬다")
    label(s, 3.9, 16, 92.3, "node를 가리지 않고 들어오는 VOC")
    vocs = [("Low VDD", "더 낮은 전압에서 동작"), ("Low Vmin", "SRAM이 버티는 최소 전압을 낮춤"),
            ("Low leakage", "대기 전력을 줄임")]
    for i, (big, sub) in enumerate(vocs):
        x = 3.8 + i * 31.8
        rect(s, x, 22, 28.8, 24, "FFFFFF", line=ACC if i == 1 else LINE, lw=1.75 if i == 1 else 0.75)
        text(s, x, 25, 28.8, 11, big, size=34, bold=True, align="c", anchor="m", color=ACC if i == 1 else INK)
        text(s, x, 37, 28.8, 6, sub, size=15, align="c", color=MUTE)
    label(s, 3.9, 52, 92.3, "그래서")
    flow(s, 3.8, 57, 92.4, ["Low VDD · Vmin · leakage VOC", "기존 cell 개선 · 신규 cell 개발",
                            "simulation에서 여러 조건 검토"], hi=2, h=10, size=16)
    message(s, 76, "검토해야 할 조건이 계속 늘고 있다")


def problem(prs, notes):
    s = frame(prs, 3, "문제: Monte Carlo 기반 Vmin 검토의 비용 한계", notes[3],
              foot="수치는 이 연구의 학습 데이터 기준 (read · write 각각)")
    label(s, 3.9, 16, 44, "조정할 수 있는 knob", color=INK, size=16)
    fit(s, TF / "s2_cell6t.png", 3.8, 21.5, 44, 40)
    for i, (dev, knobs_) in enumerate((("PU (PMOS)", "Vth · local mismatch · mobility"),
                                       ("PD · PG (NMOS)", "Vth · local mismatch · mobility"))):
        y = 63 + i * 6.5
        text(s, 3.9, y, 13, 5.4, dev, size=14, color=ACC, bold=True, anchor="m")
        text(s, 17.5, y, 30, 5.4, knobs_, size=14, color=INK, anchor="m")
    vline(s, 50.5, 16, 58, color="D9D9D9", pt=0.75)
    label(s, 53.5, 16, 42.7, "조건 하나를 검토하는 데 드는 일", color=INK, size=16)
    rows = [("조건 1개", "Monte Carlo simulation 25,000회", "VDD 5개 × VDD마다 5,000회"),
            ("knob × 값", "검토 조건이 곱으로 늘어난다", "9개 knob, 각각 여러 값 · 이 연구의 학습 조건만 1,700개"),
            ("결과", "실제로 시도할 수 있는 조건은 몇 개뿐", "충분히 비교하지 못한 채 방향을 정한다")]
    for i, (lab, big, sub) in enumerate(rows):
        y0 = 22 + i * 17
        text(s, 53.5, y0, 42, 4, lab, size=13, color=MUTE, bold=True)
        text(s, 53.5, y0 + 4, 42, 5.5, big, size=19, bold=True, color=ACC if i == 2 else INK)
        text(s, 53.5, y0 + 9.5, 42, 4, sub, size=13, color=MUTE)
        if i < 2:
            hline(s, 53.5, y0 + 15.5, 42.7, color="E6E6E6", pt=0.75)
    message(s, 80, "검토할 조건이 늘수록 Monte Carlo 비용도 그대로 늘어난다")


def proposal(prs, notes):
    s = frame(prs, 4, "제안: 순·역방향 Vmin 추정 model", notes[4],
              foot="model은 Monte Carlo simulation 결과로 한 번 학습한다")
    label(s, 20, 15.5, 22, "공정 조건 (knob 값)", align="c")
    label(s, 72, 15.5, 24.2, "Vmin", align="c")
    rect(s, 49, 21, 16, 37, "FFFFFF", line=ACC, lw=2)
    text(s, 49, 27, 16, 8, ["Vmin", "model"], size=24, color=ACC, bold=True, align="c", anchor="m")
    text(s, 49.8, 39, 14.4, 12, ["Monte Carlo", "simulation 결과로", "한 번 학습"], size=12, color=MUTE, align="c")
    rows = [("순방향 추정", ["NMOS local", "mismatch −10 %"], "Vmin = ?", "→",
             ["공정 조건이 주어지면", "→ **그 조건의 Vmin**을 찾는다"]),
            ("역방향 추정", ["NMOS local", "mismatch = ?"], f"목표 Vmin {T:.3f} V", "←",
             ["목표 Vmin이 주어지면", "→ **그 목표를 맞추는 공정 조건**을 찾는다"])]
    for i, (lab, left, right, arrow, what) in enumerate(rows):
        y0 = 21 + i * 25
        text(s, 3.8, y0, 15, 12, lab, size=17, color=ACC, bold=True, anchor="m")
        box(s, 20, y0, 22, 12, left, size=16, bold=True)
        box(s, 72, y0, 24.2, 12, right, size=17, bold=True)
        for x in (42, 65):
            text(s, x, y0, 7, 12, arrow, size=26, color=ACC, bold=True, align="c", anchor="m")
        text(s, 20, y0 + 13, 28, 9, what, size=14, color=INK, emph=ACC, spacing=1.1)
    message(s, 74, "두 방향 모두 Monte Carlo simulation을 다시 수행하지 않고 model로 추정한다")


def method(prs, notes):
    s = frame(prs, 5, "방법: 중심치·산포 기반 Vmin 추정 구조", notes[5],
              foot="Z = 중심치 / 산포 (SNMR · Vtrip 각각) · target Z는 memory density와 yield 목표로 정한다 (이 연구: 6.4)")
    steps_old = ["Vop sweep", "Monte Carlo로\nSNMR·Vtrip의\n중심치·산포", "Z =\n중심치 / 산포", "target Z를 처음\n만족하는 Vop\n= Vmin"]
    steps_new = ["Vop sweep", "model이\n중심치·산포를\n추정", "Z =\n중심치 / 산포", "target Z를 처음\n만족하는 Vop\n= Vmin"]
    label(s, 3.9, 15.5, 40, "기존 Vmin 추정")
    flow(s, 3.8, 20.5, 92.4, steps_old, size=13, h=12)
    label(s, 3.9, 36.5, 40, "이 model", color=ACC)
    flow(s, 3.8, 41.5, 92.4, steps_new, hi=1, size=13, h=12)
    seg = (92.4 - 3.2 * 3) / 4
    text(s, 3.8 + seg + 3.2, 54, seg, 4, "↑ 바뀐 건 이 단계 하나", size=13, color=ACC, bold=True, align="c")
    hline(s, 3.8, 61, 92.4, color=LINE, pt=0.75)
    label(s, 3.9, 62.5, 60, "왜 Vmin이 아니라 중심치와 산포를 학습했나", color=INK, size=16)
    why = [("①", "target Z가 바뀌어도 그대로", "memory density·yield가 달라져도\n다시 학습할 필요가 없다"),
           ("②", "더 정확하다", "중심치·산포는 조건에 따라\n부드럽게 변해 학습이 쉽다"),
           ("③", "원인이 보인다", "Vmin이 움직인 이유가\n중심치 쪽인지 산포 쪽인지 구분된다")]
    for i, (n, big, sub) in enumerate(why):
        x = 3.8 + i * 31.8
        text(s, x, 68, 4, 5, n, size=18, color=ACC, bold=True)
        text(s, x + 4, 68, 25, 5, big, size=17, color=INK, bold=True)
        text(s, x + 4, 74, 25, 9, sub.split("\n"), size=13, color=MUTE, spacing=1.1)
        if i:
            vline(s, x - 1.5, 68, 15, color="E6E6E6", pt=0.75)


def model_parts(prs, notes):
    """화면은 GP·ARD 한 줄씩. 아파트 시세 비유는 대본(구두)으로만 한다."""
    s = frame(prs, 8, "model 구성: GP와 ARD", notes[6],
              foot="GP: Gaussian Process · ARD: Automatic Relevance Determination · 자세한 구성은 논문 참고")
    cards = [("GP", "s6_gp", "수행한 조건 사이를 매끄럽게 메워", "수행하지 않은 조건을 추정한다"),
             ("ARD", "s6_ard", "knob마다 값을 흔드는 정도를", "따로 학습한다")]
    for i, (tag, fig, l1, l2) in enumerate(cards):
        x = 3.8 + i * 47.6
        text(s, x, 16, 44.6, 7, tag, size=26, color=ACC, bold=True, align="c", anchor="m")
        hline(s, x + 15, 23.5, 14.6, color=ACC, pt=1.25)
        fit(s, TF / f"{fig}.png", x + 2, 27, 40.6, 44, valign="m")
        text(s, x, 74, 44.6, 10, [l1, l2], size=18, color=INK, bold=True, align="c", spacing=1.15)
        if i:
            vline(s, x - 1.5, 16, 70, color="E6E6E6", pt=0.75)


def validation(prs, notes):
    s = frame(prs, 9, "검증: 기존 Monte Carlo 대비 Vmin 추정 정확도", notes[7],
              foot="silicon 측정이 아닌, 같은 조건에서 수행한 기존 Monte Carlo simulation 결과와 비교 · 학습에 쓰지 않은 300개 조건 중 Vmin이 VDD 범위 안에 있는 조건")
    text(s, 3.8, 15.5, 60.5, 5, ["기존 Monte Carlo simulation  vs  **제안 model**"], size=15, color=MUTE,
         emph=ACC, bold=True, align="c", anchor="m")
    fit(s, TF / "forward_accuracy.png", 3.8, 22, 60.5, 57, valign="t")
    text(s, 3.9, 80, 60, 4, "점이 대각선에 붙을수록 정확하다 (회색 띠 ±10 mV)", size=13, color=MUTE)
    blocks = [("학습에 없던 조건 · Vmin 오차 (RMSE)", ("3.98", "mV", "read"), ("5.65", "mV", "write")),
              ("학습에서 뺀 PDK corner · RMSE", ("8.5", "mV", "read"), ("5.8", "mV", "write")),
              ("Read worst · Write worst", ("FSG ✓", "", "read"), ("SFG ✓", "", "write"))]
    for i, (lab, a, b) in enumerate(blocks):
        y0 = 17 + i * 23
        text(s, 67.5, y0, 28.7, 4, lab, size=14, color=INK, bold=True)
        for j, (v, u, m) in enumerate((a, b)):
            kpi(s, 67.5 + j * 14.4, y0 + 5, 14.3, v, u, m, ACC, size=30 if u else 24)
        vline(s, 81.85, y0 + 5.5, 10, color="E6E6E6", pt=0.75)
        if i < 2:
            hline(s, 67.5, y0 + 20.5, 28.7, color="E6E6E6", pt=0.75)


def case(prs, notes):
    s = frame(prs, 10, "적용 사례: Low Vmin VOC (0.625 → 0.575 V)", notes[8],
              foot="고객명과 제품 정보는 생략 · 막대는 model이 계산한 현재 Read worst · Write worst의 Vmin")
    label(s, 3.9, 16, 43, "Low Vmin VOC")
    text(s, 3.9, 20.5, 43, 7, "“Vmin 0.625 V → 0.575 V”", size=24, color=ACC, bold=True)
    fit(s, TF / "voc_target.png", 3.8, 31, 43, 54)
    vline(s, 49.5, 16, 70, color="E6E6E6", pt=0.75)
    label(s, 52, 16, 44.2, "목표와 현재 공정의 차이", color=INK, size=16)
    text(s, 52, 22, 44.2, 7, f"목표 Vmin = 0.625 V − 0.05 V = **{T:.3f} V**", size=16, color=INK, emph=ACC, anchor="m")
    kpi(s, 52, 33, 21, f"{SM['gap_mV']['read']:.0f}", "mV", "read 부족", ACC, size=40)
    kpi(s, 75, 33, 21, f"{SM['gap_mV']['write']:.0f}", "mV", "write 부족", ACC, size=40)
    vline(s, 74, 34, 12, color="E6E6E6", pt=0.75)
    hline(s, 52, 53, 44.2, color="E6E6E6", pt=0.75)
    text(s, 52, 56, 44.2, 18, ["read와 write를 **둘 다** 맞춰야 하므로",
                               "cell Vmin = **max(read, write)**로 본다",
                               "→ 무엇을 얼마나 바꾸면 될까?"], size=17, color=INK, emph=ACC, spacing=1.35)


def knobs(prs, notes):
    s = frame(prs, 11, "knob별 Vmin 개선 효과 비교", notes[9],
              foot="나머지 조건은 고정하고 knob 하나씩 역방향 추정 · 막대 = 학습 범위 대비 바꿔야 하는 양 · 빗금 = 학습 범위 안에서는 목표에 닿지 않음")
    fit(s, TF / "lab_rank.png", 3.8, 16, 55, 76, valign="t")
    rows = [("NMOS local mismatch", NEED["l_com"], "하나만 바꿀 때 가장 작다", True),
            ("NMOS + PMOS local mismatch", "각각 " + NEED["l_com+lpu"], "둘을 같은 비율로 줄일 때", True),
            ("PMOS local mismatch", "−30 %로도 모자람", "이것만으로는 닿지 않는다", False)]
    for i, (lab, need, note, ok) in enumerate(rows):
        y0 = 17 + i * 15
        text(s, 62, y0, 34.2, 4, lab, size=13, color=MUTE, bold=True)
        text(s, 62, y0 + 4, 34.2, 5, need, size=21 if ok else 17, color=ACC if ok else MUTE, bold=True)
        text(s, 62, y0 + 9.3, 34.2, 3.5, note, size=12, color=MUTE)
        hline(s, 62, y0 + 14, 34.2, color="E6E6E6", pt=0.75)
    text(s, 62, 64, 34.2, 22, ["**Vth · mobility는 어느 것도 닿지 않는다**",
                               "read가 좋아지면 write가 나빠진다",
                               "→ 둘을 **함께 낮추는 건 local mismatch뿐**"],
         size=15, color=INK, emph=ACC, spacing=1.4)


def how_much(prs, notes):
    s = frame(prs, 12, "목표 Vmin 달성에 필요한 개선량", notes[10],
              foot="cell Vmin = max(read, write) · 별 = Read worst (FSG) · Write worst (SFG) · 나머지 7개 knob은 고정 · 학습 범위 안만 표시")
    fit(s, TF / "lab_window.png", 3.8, 16, 92.4, 62, valign="t")
    message(s, 80, f"NMOS local mismatch {NEED['l_com']}로 read · write 모두 목표 달성")


def merits(prs, notes):
    s = frame(prs, 13, "장점: 조건 변경 검토의 유연성과 학습 비용", notes[11],
              foot="표는 Read worst · Write worst가 모두 목표를 맞추는 값 (역방향 추정) · 53×는 조건 수 · VDD 수 · Monte Carlo 횟수로 계산한 비율")
    label(s, 3.9, 16, 50, "목표 Vmin을 바꿔 보면 — 필요한 local mismatch 감소량", color=INK, size=16)
    head = ("목표 Vmin", "NMOS만", "NMOS + PMOS 각각")
    rows = [(f"{r['target']:.3f} V" + (" (VOC)" if abs(r["target"] - T) < 1e-9 else ""), r["l_com"], r["l_com+lpu"])
            for r in SM["target_sweep"]]
    xs, ws = (3.8, 20, 36), (16, 16, 19)
    for j, h in enumerate(head):
        text(s, xs[j], 22.5, ws[j], 5, h, size=14, color=MUTE, bold=True, align="c", anchor="m")
    hline(s, 3.8, 28, 51.2, color=INK, pt=1)
    for i, row in enumerate(rows):
        y0 = 29 + i * 10
        voc = row[0].endswith("(VOC)")
        if voc:
            rect(s, 3.8, y0, 51.2, 10, SOFT)
        for j, v in enumerate(row):
            text(s, xs[j], y0, ws[j], 10, v, size=20 if j else 17, bold=j > 0 or voc,
                 color=(ACC if voc else INK) if j else (INK if voc else MUTE), align="c", anchor="m")
        hline(s, 3.8, y0 + 10, 51.2, color="E6E6E6", pt=0.75)
    text(s, 3.9, 71, 51, 4, "목표를 바꿀 때마다 Monte Carlo simulation이 더 들지 않는다", size=14, color=INK, bold=True)
    vline(s, 59, 16, 62, color="E6E6E6", pt=0.75)
    label(s, 62, 16, 34.2, "처음 학습 비용도 줄일 수 있다", color=INK, size=16)
    for i, (lab, before, after) in enumerate((("학습 조건", "1,700", "400"), ("VDD 수", "5", "4"),
                                              ("Monte Carlo 횟수", "5,000", "500"))):
        y0 = 23 + i * 9
        hline(s, 62, y0, 34.2, color="E6E6E6", pt=0.75)
        text(s, 62, y0 + 0.8, 14, 7, lab, size=13, color=MUTE, anchor="m")
        text(s, 76, y0 + 0.8, 20.2, 7, f"{before}  →  **{after}**", size=18, color=INK, emph=INK, anchor="m")
    hline(s, 62, 50, 34.2, color="E6E6E6", pt=0.75)
    kpi(s, 62, 53, 17, "1/53", "", "학습 비용 (read)", ACC, size=30)
    kpi(s, 79.2, 53, 17, "+3.5", "mV", "늘어나는 오차", ACC, size=30)
    vline(s, 79.1, 54, 10, color="E6E6E6", pt=0.75)
    message(s, 80, "검토 조건이 늘어도 Monte Carlo 비용은 늘지 않는다")


def summary(prs, notes):
    s = frame(prs, 14, "정리: model 기반 Vmin 검토의 의의", notes[12])
    cards = [("순·역방향 추정 model", ["Monte Carlo simulation 결과로", "Vmin을 두 방향으로 추정", "오차 read 4 mV · write 6 mV"]),
             ("효과적인 knob 도출", ["read · write를 함께 낮추는 건", "local mismatch", f"Vmin −50 mV ← NMOS {NEED['l_com']}"]),
             ("마음껏 검토", ["목표 · knob 조합을 바꿔도", "Monte Carlo simulation 추가 없음", "학습 비용도 1/53까지"])]
    for i, (head_, body) in enumerate(cards):
        x = 3.8 + i * 31.8
        text(s, x, 17, 28.8, 7, head_, size=20, color=ACC, bold=True, align="c", anchor="m")
        hline(s, x + 9, 25, 10.8, color=ACC, pt=1.25)
        text(s, x, 30, 28.8, 22, body, size=17, color=INK, align="c", spacing=1.6)
        if i:
            vline(s, x - 1.5, 17, 37, color="E6E6E6", pt=0.75)
    message(s, 60, "model로 먼저 좁히고, Monte Carlo simulation은 확인에만 쓴다")
    text(s, 3.9, 78, 60, 6, "한계 — 결과는 나머지 조건을 고정했을 때의 값 · 비교 기준은 Monte Carlo simulation",
         size=12, color=MUTE)
    text(s, 66, 75, 30.2, 9, "감사합니다", size=28, color=ACC, bold=True, align="r", anchor="m")


PHYS_NOTE = ("이 framework를 physics-guided라고 부르는 이유는, 물리로 정해진 부분은 학습시키지 않았기 때문입니다. "
             "먼저 Vmin 계산식입니다. 중심치를 산포로 나눈 z가 target Z를 처음 넘는 전압이 Vmin인데, "
             "target Z는 array 크기와 yield 목표로 계산한 값이고, 이 식에는 학습하는 값이 하나도 없습니다. "
             "둘째, 이 계산은 공급 전압에 대해 단조라서, 목표 Vmin에서 공정 조건을 거꾸로 찾는 역방향 추정이 이분 탐색으로 안정적으로 풀립니다. "
             "셋째, 데이터도 물리로 점검했습니다. 공급 전압을 올리면 margin 중심치는 커져야 하므로, 이를 어기는 값은 전사 오류로 보고 바로잡았습니다. "
             "마지막으로 GP는 공정 조건이 조금 바뀌면 margin도 조금 바뀐다는 연속성을 전제로 하고, 산포는 Monte Carlo 오차가 일정해지는 log 축에서 학습했습니다.")


def physics(prs, notes, n=6):
    """physics-guided의 실체. 근거: 최종 논문(make_docx.js s1c·s2a·s2b, paper_kr_D §II-A·§IV) + v_b_forward.py.
    논문 D판의 단조성 패널티 · 코너 앵커링은 최종 model에 없으므로 넣지 않는다."""
    s = frame(prs, n, "physics-guided 요소: 학습하지 않는 Vmin 계산과 단조성", notes.get(n, PHYS_NOTE),
              foot="p_fail = −ln(0.99) / (128 × 10⁶) · 128 Mb array, 99 % yield · 데이터 점검 기준: 위반이 Monte Carlo 표준오차의 3배 초과")
    label(s, 3.9, 16, 40, "Vmin 계산식 (physics layer)", color=INK, size=16)
    rect(s, 3.8, 22, 40, 34, "FFFFFF", line=ACC, lw=1.75)
    text(s, 5.5, 24, 36.6, 30, ["z(Vop) = 중심치 / 산포", "",
                                "target Z = Φ⁻¹(1 − p_fail) = **6.4**", "",
                                "Vmin = z가 target Z를", "처음 넘는 가장 낮은 Vop"],
         size=17, color=INK, emph=ACC, align="c", spacing=1.15)
    text(s, 3.8, 58, 40, 10, ["학습하는 값이 **하나도 없다**",
                              "→ model은 중심치 · 산포만 배운다"], size=15, color=INK, emph=ACC, align="c", spacing=1.3)
    vline(s, 47.5, 16, 66, color="E6E6E6", pt=0.75)
    rows = [("①", "계산식은 학습하지 않는다", "데이터가 적은 곳에서도 Vmin의 정의는 흔들리지 않고,\ntarget Z가 바뀌어도 다시 학습할 필요가 없다"),
            ("②", "VDD에 대한 단조성", "계산이 VDD에 단조 → 역방향 추정은 이분 탐색으로 안정적으로\n풀린다 (단조 구간을 먼저 확인한 뒤 사용)"),
            ("③", "데이터의 단조성 점검", "VDD를 올리면 margin 중심치는 커져야 한다 → 어긋난 값은\n전사 오류로 보고 정정 (read 31셀, write 12셀)"),
            ("④", "공정 연속성", "GP: 공정 조건이 조금 바뀌면 margin도 조금 바뀐다\n산포는 Monte Carlo 오차가 일정해지는 log 축에서 학습")]
    for i, (num, head_, body) in enumerate(rows):
        y0 = 16 + i * 17
        text(s, 50.5, y0, 4, 5, num, size=18, color=ACC, bold=True)
        text(s, 54.5, y0, 41.7, 5, head_, size=17, color=INK, bold=True)
        text(s, 54.5, y0 + 5.5, 41.7, 9, body.split("\n"), size=13, color=MUTE, spacing=1.15)
        if i < 3:
            hline(s, 50.5, y0 + 15.5, 45.7, color="E6E6E6", pt=0.75)
    message(s, 85, "물리로 정해진 부분은 식으로 두고, 모르는 부분(중심치 · 산포)만 학습한다")


AXES_NOTE = ("입력 축은 Vth, local mismatch, mobility 세 종류를 NMOS와 PMOS에 나눠 잡았습니다. "
             "NMOS는 pull-down과 pass-gate 두 개가 있는데, 같은 NMOS라 대부분 함께 움직입니다. "
             "그래서 기본은 둘을 함께 움직이는 공통 축으로 두고, 완전히 같지는 않으니 둘의 차이를 skew 축 하나로 허용했습니다. "
             "범위는 Vth가 ±60 mV로 PDK corner를 모두 안에 포함하고, local mismatch와 mobility는 ±30 %로 "
             "개선 목표의 경계가 범위 안에 들어오도록 넉넉하게 잡았습니다. skew는 공통 범위의 1/4에서 1/3 정도로 작게 두었습니다.")


def axes(prs, notes, n=7):
    """초안: 입력 9개 축을 고른 이유와 범위. 근거는 condition_gen.py v2.1 · docs/decisions/legacy_design_audit_20260714.md §5."""
    s = frame(prs, n, "입력 축: NMOS 공통 · PD–PG skew · PMOS 구성과 범위", notes.get(n, AXES_NOTE),
              foot="범위 근거: 조건 생성기 v2.1 설계 검토 (2026-07-14) · PDK corner = FFG · FSG · SFG · SSG의 NMOS · PMOS Vth shift")
    label(s, 3.9, 16, 52, "입력 9개 축과 범위", color=INK, size=16)
    cols = [("NMOS 공통", "PD · PG 함께"), ("PD–PG skew", "둘의 차이"), ("PMOS", "PU")]
    rows = [("Vth shift", ("±60 mV", "±20 mV", "±60 mV")),
            ("local mismatch", ("×0.7 ~ 1.3", "±0.075", "×0.7 ~ 1.3")),
            ("mobility", ("×0.7 ~ 1.3", "±0.075", "×0.7 ~ 1.3"))]
    x0, cw = 17.5, 12.6
    rect(s, x0 + cw, 21, cw, 37.5, "", line=ACC, lw=1.25)        # skew 열 강조 (채우지 않음, 글씨보다 먼저)
    for j, (h1, h2) in enumerate(cols):
        x = x0 + j * cw
        text(s, x, 22, cw, 4, h1, size=14, color=ACC if j == 1 else INK, bold=True, align="c")
        text(s, x, 26, cw, 3.5, h2, size=11, color=MUTE, align="c")
    hline(s, 3.8, 30.5, 51.5, color=INK, pt=1)
    for i, (name, vals) in enumerate(rows):
        y0 = 31 + i * 9
        text(s, 3.9, y0, 13.5, 9, name, size=14, color=INK, bold=True, anchor="m")
        for j, v in enumerate(vals):
            text(s, x0 + j * cw, y0, cw, 9, v, size=16, color=ACC if j == 1 else INK, bold=j == 1,
                 align="c", anchor="m")
        hline(s, 3.8, y0 + 9, 51.5, color="E6E6E6", pt=0.75)
    box(s, 3.8, 63, 51.5, 9, ["PG = 공통 + skew      PD = 공통 − skew"], size=16, bold=True)
    text(s, 3.9, 73.5, 51.4, 4, "local mismatch · mobility도 같은 방식 (배율 = 공통 ± skew)", size=12, color=MUTE)
    text(s, 3.9, 79, 51.4, 8, ["**표본 선정**: read · write 각 2,000 조건 · 각 축은 범위 안에서 random (PCG64)",
                               "Vmin이 나빠지는 사분면에 가중 — read FSG 45 %, write SFG 45 %"],
         size=13, color=INK, emph=ACC, spacing=1.25)
    vline(s, 58, 16, 66, color="E6E6E6", pt=0.75)
    blocks = [("왜 공통 + skew인가", ["PD · PG는 같은 NMOS라 **대부분 함께** 움직인다",
                                    "완전히 같지는 않으니 **차이는 skew 축 하나**로 허용",
                                    "따로 뽑으면 PG −30 %, PD +30 % 같은", "비현실적인 조건에 표본을 낭비한다"]),
              ("공통 축의 범위", ["Vth **±60 mV**: PDK corner (±29~45 mV)를 모두 포함",
                               "local mismatch · mobility **±30 %**:",
                               "보통 corner 폭 (±10~15 %)보다 넉넉하게",
                               "→ 개선 경계가 범위 안에 들어온다"]),
              ("skew 축의 범위", ["공통 범위의 **1/4 ~ 1/3** (Vth ±20 mV, 배율 ±0.075)",
                               "→ PD · PG가 함께 움직이는 정도(상관) 약 0.8 ~ 0.9"])]
    y = 16
    for head_, body in blocks:
        text(s, 61, y, 35.2, 4.5, head_, size=15, color=ACC, bold=True)
        text(s, 61, y + 5, 35.2, 4.6 * len(body), body, size=13, color=INK, emph=INK, spacing=1.15)
        y += 5 + 4.6 * len(body) + 3


def cost_compare(prs, notes, n=12):
    """기존 crossed DOE와 GP 학습의 deck 수 비교 (deck = 조건 1개 × VDD 1개의 MC run)."""
    s = frame(prs, n, "Simulation 비용: 기존 DOE와 비교", notes.get(n, ""),
              foot="deck = 조건 1개 × VDD 1개의 Monte Carlo run · 조건당 Monte Carlo 횟수는 같다고 가정")
    cols = (("기존: 9D crossed DOE", 3.8, False,
             [("2D contour", "CCD 15점 × 지표 2 × VDD 5", "150"),
              ("나머지 7개 knob", "3-level grid 3⁷", "× 2,187")], "328,050"),
            ("제안: GP 학습", 53, True,
             [("SNMR · BWRM", "조건 2,000 × 지표 2", "4,000"),
              ("VDD", "5개", "× 5")], "20,000"))
    for title, x, hi, rows, total in cols:
        label(s, x, 16, 43.2, title, color=ACC if hi else INK, size=17)
        hline(s, x, 21.5, 43.2, color=INK, pt=1)
        for i, (lab, body, val) in enumerate(rows):
            y = 23 + i * 10
            text(s, x, y, 14, 9, lab, size=14, color=MUTE, bold=True, anchor="m")
            text(s, x + 14, y, 21, 9, body, size=15, color=INK, anchor="m")
            text(s, x + 35, y, 8.2, 9, val, size=18, color=INK, align="r", anchor="m")
            hline(s, x, y + 9.5, 43.2, color="E6E6E6", pt=0.75)
        text(s, x, 46, 15, 12, "총 deck", size=15, color=INK, bold=True, anchor="m")
        text(s, x + 15, 46, 28.2, 12, total, size=34, color=ACC if hi else INK, bold=True,
             align="r", anchor="m")
    vline(s, 50, 16, 44, color="D9D9D9", pt=0.75)
    kpi(s, 30, 63, 40, "1/16", "", "deck 수 (20,000 ÷ 328,050)", ACC, size=36)
    message(s, 80, "9개 knob 전체를 보는 데 필요한 simulation이 1/16로 줄어든다")


def build(notes) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = d.W, d.H
    for f in (cover, background, problem, proposal, method, physics, axes, model_parts,
              validation, case, knobs, how_much, merits, summary):
        f(prs, notes)
    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.has_text_frame and shp.text_frame.text == "SRAM Vmin surrogate":
                shp.text_frame.paragraphs[0].runs[0].text = BRAND
    return prs


if __name__ == "__main__":
    if "--draft-cost" in sys.argv:       # 비용 비교 장표 1장만 따로 만든다 (본 deck에는 아직 넣지 않음)
        prs = Presentation()
        prs.slide_width, prs.slide_height = d.W, d.H
        cost_compare(prs, {})
        for shp in prs.slides[0].shapes:
            if shp.has_text_frame and shp.text_frame.text == "SRAM Vmin surrogate":
                shp.text_frame.paragraphs[0].runs[0].text = BRAND
        out = HERE / "draft_cost.pptx"
        prs.save(out)
        print("saved", out)
        raise SystemExit
    if "--draft-axes" in sys.argv:       # 입력 축 장표 초안 1장만 따로 만든다 (본 deck에는 아직 넣지 않음)
        prs = Presentation()
        prs.slide_width, prs.slide_height = d.W, d.H
        axes(prs, {})
        for shp in prs.slides[0].shapes:
            if shp.has_text_frame and shp.text_frame.text == "SRAM Vmin surrogate":
                shp.text_frame.paragraphs[0].runs[0].text = BRAND
        out = HERE / "draft_axes.pptx"
        prs.save(out)
        print("saved", out)
        raise SystemExit
    notes = load_notes()
    build(notes).save(OUT)
    chars = sum(len(re.sub(r"\s", "", t)) for t in notes.values())
    print(f"saved {OUT}  ·  {len(notes)}장 · 노트 대본({SCRIPT.name}) 공백 제외 {chars}자 ≈ {chars / 330:.1f}분")
